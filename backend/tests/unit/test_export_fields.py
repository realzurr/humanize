"""Testes unitários dos campos dinâmicos OpenXML (Word) da Onda 6."""

from __future__ import annotations

import io

import docx
import pytest

from app.services.document_export.fields import (
    FieldKind,
    add_bookmark,
    add_numpages_field,
    add_page_field,
    add_ref_field,
    add_seq_field,
    add_styleref_field,
    add_toc_field,
    enable_update_fields_on_open,
    extract_field_instructions,
)


def _save_to_bytes(doc: docx.document.Document) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_all_six_word_fields_are_generated_and_extracted() -> None:
    doc = docx.Document()
    enable_update_fields_on_open(doc)

    p_toc = doc.add_paragraph()
    add_toc_field(p_toc, levels="1-3")

    p_head = doc.add_paragraph("1 INTRODUÇÃO")
    add_bookmark(p_head, "sec_intro", 1)

    p_body = doc.add_paragraph("Conforme a seção ")
    add_ref_field(p_body, "sec_intro", hyperlink=True)

    p_seq = doc.add_paragraph("Tabela ")
    add_seq_field(p_seq, "Tabela")

    header_p = doc.sections[0].header.paragraphs[0]
    add_styleref_field(header_p, "Heading 1")
    add_page_field(header_p)
    add_numpages_field(header_p)

    raw = _save_to_bytes(doc)
    instrucoes = extract_field_instructions(raw)

    assert any(i.startswith('TOC \\o "1-3"') for i in instrucoes)
    assert any(i == "PAGE" for i in instrucoes)
    assert any(i == "NUMPAGES" for i in instrucoes)
    assert any(i == "SEQ Tabela \\* ARABIC" for i in instrucoes)
    assert any(i == 'STYLEREF "Heading 1"' for i in instrucoes)
    assert any(i == "REF sec_intro \\h" for i in instrucoes)
    assert FieldKind.TOC == "TOC"


def test_invalid_bookmark_raises_value_error() -> None:
    doc = docx.Document()
    p = doc.add_paragraph("Teste")
    with pytest.raises(ValueError, match="marcador inválido"):
        add_bookmark(p, "!!!", 1)
    with pytest.raises(ValueError, match="marcador inválido"):
        add_ref_field(p, "   ")


def test_enable_update_fields_idempotent() -> None:
    doc = docx.Document()
    enable_update_fields_on_open(doc)
    enable_update_fields_on_open(doc)
    raw = _save_to_bytes(doc)
    assert len(raw) > 0
