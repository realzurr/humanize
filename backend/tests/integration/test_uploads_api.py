"""Testes de integração do upload seguro de documentos."""

from __future__ import annotations

from fastapi.testclient import TestClient


def test_upload_texto_valido(app_with_db) -> None:
    with TestClient(app_with_db) as client:
        resposta = client.post(
            "/api/v1/uploads",
            files={"arquivo": ("documento.txt", b"conteudo de teste", "text/plain")},
        )
        assert resposta.status_code == 201
        corpo = resposta.json()
        assert corpo["format"] == "text"
        assert corpo["texto"] == "conteudo de teste"
        assert corpo["filename"] == "documento.txt"


def test_upload_rejeita_extensao_nao_suportada(app_with_db) -> None:
    with TestClient(app_with_db) as client:
        resposta = client.post(
            "/api/v1/uploads",
            files={"arquivo": ("programa.exe", b"MZ\x00\x00\x00", "application/octet-stream")},
        )
        assert resposta.status_code == 422


def test_upload_rejeita_docx_invalido(app_with_db) -> None:
    with TestClient(app_with_db) as client:
        resposta = client.post(
            "/api/v1/uploads",
            files={"arquivo": ("falso.docx", b"nao-eh-um-zip", "application/octet-stream")},
        )
        assert resposta.status_code == 422


def test_upload_rejeita_tamanho_excedido(app_with_db) -> None:
    app_with_db.state.settings.max_upload_bytes = 10
    with TestClient(app_with_db) as client:
        resposta = client.post(
            "/api/v1/uploads",
            files={"arquivo": ("grande.txt", b"x" * 100, "text/plain")},
        )
        assert resposta.status_code == 413


def test_upload_docx_extrai_paragrafos(app_with_db) -> None:
    """GAP-05: um DOCX real é lido e tem seus parágrafos extraídos."""
    from io import BytesIO

    from docx import Document

    buffer = BytesIO()
    documento = Document()
    documento.add_paragraph("Primeiro parágrafo")
    documento.add_paragraph("Segundo parágrafo")
    documento.save(buffer)

    with TestClient(app_with_db) as client:
        resposta = client.post(
            "/api/v1/uploads",
            files={
                "arquivo": (
                    "trabalho.docx",
                    buffer.getvalue(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        assert resposta.status_code == 201
        corpo = resposta.json()
        assert corpo["format"] == "docx"
        assert "Primeiro parágrafo" in corpo["texto"]


def test_upload_mesmo_nome_nao_sobrescreve(app_with_db, tmp_path) -> None:
    """GAP-05: uploads de mesmo nome são armazenados com nomes únicos em disco."""
    app_with_db.state.settings.storage_dir = tmp_path
    with TestClient(app_with_db) as client:
        primeiro = client.post(
            "/api/v1/uploads", files={"arquivo": ("artigo.txt", b"conteudo um", "text/plain")}
        )
        segundo = client.post(
            "/api/v1/uploads", files={"arquivo": ("artigo.txt", b"conteudo dois", "text/plain")}
        )
        assert primeiro.status_code == 201
        assert segundo.status_code == 201
        assert primeiro.json()["id"] != segundo.json()["id"]

        arquivos = list((tmp_path / "uploads" / "local").iterdir())
        assert len(arquivos) == 2
