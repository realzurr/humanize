"""Construção e inspeção de campos dinâmicos OpenXML (Word).

Materializa campos reais do Word (``TOC``, ``PAGE``, ``NUMPAGES``, ``SEQ``,
``STYLEREF``, ``REF``) na árvore XML do ``python-docx``, evitando texto fixo
simulado (ver `ADR-0012`). O sumário e as referências cruzadas dependem de
atualização de campos ao abrir o documento no editor, por isso habilitamos
``w:updateFields`` em ``settings.xml``.
"""

from __future__ import annotations

import io
import xml.etree.ElementTree as ET
import zipfile
from enum import StrEnum
from typing import TYPE_CHECKING

from docx.oxml import OxmlElement
from docx.oxml.ns import qn

if TYPE_CHECKING:
    from docx.document import Document as DocumentObject
    from docx.text.paragraph import Paragraph

_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


class FieldKind(StrEnum):
    """Tipos de campo dinâmico suportados na exportação DOCX."""

    TOC = "TOC"
    PAGE = "PAGE"
    NUMPAGES = "NUMPAGES"
    SEQ = "SEQ"
    STYLEREF = "STYLEREF"
    REF = "REF"


def add_field(
    paragraph: Paragraph,
    instruction: str,
    *,
    placeholder: str = "",
    dirty: bool = False,
) -> None:
    """Insere um campo complexo OpenXML (``begin`` → ``instrText`` → ``separate`` → ``end``)."""
    run_begin = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    if dirty:
        fld_begin.set(qn("w:dirty"), "true")
    run_begin._r.append(fld_begin)

    run_instr = paragraph.add_run()
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = f" {instruction.strip()} "
    run_instr._r.append(instr_text)

    run_sep = paragraph.add_run()
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    run_sep._r.append(fld_sep)

    if placeholder:
        paragraph.add_run(placeholder)

    run_end = paragraph.add_run()
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run_end._r.append(fld_end)


def add_toc_field(
    paragraph: Paragraph,
    *,
    levels: str = "1-3",
    placeholder: str = (
        "Sumário automático (atualize os campos no Word ou LibreOffice — F9 — para paginar)"
    ),
) -> None:
    """Insere o campo de sumário automático real (``TOC``)."""
    instruction = f'TOC \\o "{levels}" \\h \\z \\u'
    add_field(paragraph, instruction, placeholder=placeholder, dirty=True)


def add_page_field(paragraph: Paragraph) -> None:
    """Insere o campo de número da página atual (``PAGE``)."""
    add_field(paragraph, "PAGE", placeholder="1")


def add_numpages_field(paragraph: Paragraph) -> None:
    """Insere o campo de total de páginas do documento (``NUMPAGES``)."""
    add_field(paragraph, "NUMPAGES", placeholder="1")


def add_seq_field(paragraph: Paragraph, identifier: str = "Tabela") -> None:
    """Insere o campo de sequência numérica automática (``SEQ``)."""
    ident_limpo = "".join(ch for ch in identifier if ch.isalnum() or ch == "_") or "Item"
    add_field(paragraph, f"SEQ {ident_limpo} \\* ARABIC", placeholder="1", dirty=True)


def add_styleref_field(paragraph: Paragraph, style_name: str = "Heading 1") -> None:
    """Insere referência ao texto do estilo corrente no cabeçalho/rodapé (``STYLEREF``)."""
    estilo_escapado = style_name.replace('"', "")
    add_field(paragraph, f'STYLEREF "{estilo_escapado}"', dirty=True)


def add_ref_field(
    paragraph: Paragraph,
    bookmark: str,
    *,
    hyperlink: bool = True,
    placeholder: str = "",
) -> None:
    """Insere referência cruzada a um marcador (``REF``)."""
    marcador = "".join(ch for ch in bookmark if ch.isalnum() or ch in ("_", "-"))
    if not marcador:
        raise ValueError("nome de marcador inválido para campo REF")
    flags = " \\h" if hyperlink else ""
    add_field(
        paragraph,
        f"REF {marcador}{flags}",
        placeholder=placeholder or marcador,
        dirty=True,
    )


def add_bookmark(paragraph: Paragraph, name: str, bookmark_id: int) -> None:
    """Envolve o conteúdo do parágrafo em um marcador (``w:bookmarkStart``/``End``)."""
    marcador = "".join(ch for ch in name if ch.isalnum() or ch in ("_", "-"))
    if not marcador:
        raise ValueError("nome de marcador inválido")
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), marcador)

    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))

    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def enable_update_fields_on_open(document: DocumentObject) -> None:
    """Configura ``w:updateFields`` em ``settings.xml`` para atualizar campos ao abrir."""
    settings_el = document.settings.element
    existing = settings_el.find(qn("w:updateFields"))
    if existing is None:
        update_fields = OxmlElement("w:updateFields")
        update_fields.set(qn("w:val"), "true")
        settings_el.append(update_fields)
    else:
        existing.set(qn("w:val"), "true")


def extract_field_instructions(docx_bytes: bytes) -> list[str]:
    """Extrai todas as instruções ``w:instrText`` do documento, cabeçalhos e rodapés."""
    instrucoes: list[str] = []
    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as archive:
        for nome in sorted(archive.namelist()):
            if not (
                nome == "word/document.xml"
                or (nome.startswith("word/header") and nome.endswith(".xml"))
                or (nome.startswith("word/footer") and nome.endswith(".xml"))
            ):
                continue
            xml_bytes = archive.read(nome)
            # XML interno de pacote DOCX gerado/validado localmente.
            root = ET.fromstring(xml_bytes)  # noqa: S314
            for elem in root.iter(f"{{{_W_NS}}}instrText"):
                if elem.text and elem.text.strip():
                    instrucoes.append(elem.text.strip())
    return instrucoes
